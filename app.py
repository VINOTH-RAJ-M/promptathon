"""Promptathon: admin + participant web system (Flask + Socket.IO).
Run:  python app.py   ->  admin: /admin   participants: /
"""
import os, io, csv, json, time, socket
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from flask import Flask, render_template, request, redirect, session, Response
from flask_socketio import SocketIO, emit, join_room

BASE = os.path.dirname(os.path.abspath(__file__))
ADMIN_PASSWORD = os.environ.get("ADMIN_PASSWORD", "DSUAIML")
MAX_WARNINGS = int(os.environ.get("MAX_WARNINGS", "2"))
STATE_FILE = os.path.join(BASE, "state.json")

app = Flask(__name__)
app.secret_key = os.environ.get("SECRET_KEY", "change-me-promptathon")
socketio = SocketIO(app, async_mode="threading", cors_allowed_origins="*")

with open(os.path.join(BASE, "rounds.json"), encoding="utf-8") as f:
    ROUNDS = json.load(f)

state = {"current": None, "status": "idle", "end_time": None,
         "participants": {}, "submissions": {}, "scores": {}}
if os.path.exists(STATE_FILE):
    with open(STATE_FILE, encoding="utf-8") as f:
        state.update(json.load(f))

ONLINE = {}        # roll -> set of socket ids
SID_ROLL = {}      # socket id -> roll
ADMIN_SIDS = set()


import threading
STATE_LOCK = threading.Lock()

def save():
    with STATE_LOCK:
        try:
            temp_file = STATE_FILE + ".tmp"
            with open(temp_file, "w", encoding="utf-8") as f:
                json.dump(state, f, ensure_ascii=False, indent=1)
            os.replace(temp_file, STATE_FILE)
        except Exception:
            try:
                with open(STATE_FILE, "w", encoding="utf-8") as f:
                    json.dump(state, f, ensure_ascii=False, indent=1)
            except Exception:
                pass


import re, difflib

STOP_WORDS = {"the", "a", "an", "in", "on", "of", "and", "or", "to", "for", "with", "is", "are", "as", "at", "by", "this", "that", "it", "from", "be", "your", "you", "i", "we", "our"}

def tokenize(text):
    words = re.findall(r'\b[a-zA-Z0-9_\-\$]+\b', str(text).lower())
    return [w for w in words if w not in STOP_WORDS]

def evaluate_theme_score(user_text, key_answer, max_marks):
    if not user_text or not str(user_text).strip():
        return 0.0
    u_text = str(user_text).strip()
    k_text = str(key_answer).strip()
    
    u_tokens = set(tokenize(u_text))
    k_tokens = set(tokenize(k_text))
    
    if not k_tokens:
        return float(max_marks)
    
    # 1. Token overlap (Keyword & concept coverage)
    overlap = len(u_tokens & k_tokens)
    token_coverage = min(1.0, overlap / max(1, len(k_tokens) * 0.6))
    
    # 2. Sequence similarity
    seq_ratio = difflib.SequenceMatcher(None, u_text.lower(), k_text.lower()).ratio()
    
    # 3. Completeness / length ratio
    len_ratio = min(1.0, len(u_text) / max(30, len(k_text) * 0.5))
    
    # Combined moderate similarity score (0.0 to 1.0)
    sim = (0.50 * token_coverage) + (0.30 * seq_ratio) + (0.20 * len_ratio)
    
    # Moderate grading curve for 1st year participants
    if sim >= 0.50:
        multiplier = 0.90 + 0.10 * min(1.0, (sim - 0.50) / 0.40)
    elif sim >= 0.32:
        multiplier = 0.75 + 0.15 * ((sim - 0.32) / 0.18)
    elif sim >= 0.18:
        multiplier = 0.55 + 0.20 * ((sim - 0.18) / 0.14)
    elif sim >= 0.08:
        multiplier = 0.35 + 0.20 * ((sim - 0.08) / 0.10)
    else:
        multiplier = max(0.15, sim * 2.5)
        
    raw_score = multiplier * max_marks
    return round(min(float(max_marks), max(0.0, raw_score)), 1)


def auto_evaluate_round(round_idx, answers, prompt_text):
    if not (0 <= round_idx < len(ROUNDS)):
        return 0.0
    r_config = ROUNDS[round_idx]
    questions = r_config.get("questions", [])
    max_round_marks = float(r_config.get("max_marks", 40))
    theme_marks = float(r_config.get("theme_marks", 10))
    
    total = 0.0
    if isinstance(answers, dict) and len(answers) > 0:
        for k, ans in answers.items():
            ans_str = str(ans).strip()
            if ans_str:
                idx = int(k) if str(k).isdigit() else 0
                if idx < len(questions):
                    q = questions[idx]
                    q_max = float(q.get("max_marks", theme_marks))
                    q_key = q.get("key_answer", "")
                    total += evaluate_theme_score(ans_str, q_key, q_max)
    elif prompt_text:
        # If single combined text prompt
        best_score = 0.0
        for q in questions:
            q_max = float(q.get("max_marks", theme_marks))
            q_key = q.get("key_answer", "")
            s = evaluate_theme_score(prompt_text, q_key, q_max)
            if s > best_score:
                best_score = s
        total = best_score
        
    return round(min(max_round_marks, total), 1)


def round_payload():
    cur, st = state["current"], state["status"]
    now = time.time()
    p = {"status": st, "index": cur, "total": len(ROUNDS),
         "start_time": state.get("start_time"),
         "end_time": state.get("end_time"), "server_now": now}
    if cur is not None and st in ("running", "ended"):
        p["title"] = ROUNDS[cur]["title"]
        p["round_no"] = cur + 1
        p["total_minutes"] = ROUNDS[cur].get("minutes", 40)
        p["theme_minutes"] = ROUNDS[cur].get("theme_minutes", 10)
        if st in ("running", "ended"):
            p["problem"] = ROUNDS[cur].get("problem", "")
            p["instruction"] = ROUNDS[cur].get("instruction", "")
            
            # Calculate current active slot based on elapsed time
            start = state.get("start_time", now)
            elapsed = max(0, now - start)
            theme_sec = p["theme_minutes"] * 60
            all_questions = ROUNDS[cur].get("questions", [])
            active_slot = min(len(all_questions) - 1, int(elapsed // theme_sec)) if all_questions else 0
            p["active_slot"] = active_slot
            
            # Send questions: reveal only up to active_slot; future questions are locked and content is hidden
            safe_questions = []
            for idx, q in enumerate(all_questions):
                q_mins = q.get("minutes", p["theme_minutes"])
                if st == "ended" or idx <= active_slot:
                    safe_questions.append({
                        "id": q.get("id"),
                        "title": q.get("title"),
                        "minutes": q_mins,
                        "details": q.get("details", ""),
                        "unlocked": True,
                        "slot": idx
                    })
                else:
                    safe_questions.append({
                        "id": q.get("id"),
                        "title": f"Theme {idx+1} (Locked)",
                        "minutes": q_mins,
                        "details": "🔒 This question is locked and will be revealed once the current theme time completes.",
                        "unlocked": False,
                        "slot": idx
                    })
            p["questions"] = safe_questions
    return p


def me_payload(roll):
    pr = state["participants"].get(roll, {})
    sub = None
    if state["current"] is not None:
        sub = state["submissions"].get(f"{roll}|{state['current']}")
    return {"warnings": pr.get("warnings", 0), "max": MAX_WARNINGS,
            "disqualified": pr.get("disqualified", False),
            "disqualified_reason": pr.get("disqualified_reason", ""),
            "is_submitted": bool(sub and sub.get("submitted")),
            "prompt": sub["prompt"] if sub else "",
            "answers": sub.get("answers", {}) if sub else {}}



def admin_data():
    people = []
    for roll, pr in state["participants"].items():
        total = sum(float(v) for k, v in state["scores"].items() if k.startswith(roll + "|") and str(v).replace('.', '', 1).isdigit())
        done = sum(1 for k in state["submissions"] if k.startswith(roll + "|"))
        people.append({"roll": roll, "name": pr["name"], "warnings": pr.get("warnings", 0),
                       "disqualified": pr.get("disqualified", False),
                       "last_event": pr.get("last_event", ""), "online": bool(ONLINE.get(roll)),
                       "submitted": done, "total": round(total, 1)})
    people.sort(key=lambda x: (-x["total"], x["name"].lower()))
    subs = []
    for key, s in state["submissions"].items():
        subs.append(dict(s, key=key, score=state["scores"].get(key, "")))
    subs.sort(key=lambda x: x["time"])
    return {"round": round_payload(), "rounds": ROUNDS, "people": people,
            "subs": subs, "max": MAX_WARNINGS}


def push():
    socketio.emit("round_state", round_payload(), to="participants")
    socketio.emit("admin_data", admin_data(), to="admins")


def push_admin():
    socketio.emit("admin_data", admin_data(), to="admins")


def ticker():
    while True:
        socketio.sleep(1)
        if state["status"] == "running":
            now = time.time()
            if now >= state["end_time"]:
                state["status"] = "ended"
                save()
                push()
            else:
                cur = state.get("current")
                if cur is not None and 0 <= cur < len(ROUNDS):
                    q_list = ROUNDS[cur].get("questions", [])
                    theme_mins = ROUNDS[cur].get("theme_minutes", 10)
                    start = state.get("start_time", now)
                    elapsed = max(0, now - start)
                    cur_slot = min(len(q_list) - 1, int(elapsed // (theme_mins * 60))) if q_list else 0
                    if cur_slot != state.get("active_slot"):
                        state["active_slot"] = cur_slot
                        save()
                        push()


ticker_started = False
def start_ticker_if_needed():
    global ticker_started
    if not ticker_started:
        socketio.start_background_task(ticker)
        ticker_started = True

@app.before_request
def before_request_hook():
    start_ticker_if_needed()


# ---------- pages ----------
@app.route("/")
def index():
    return render_template("index.html", error=request.args.get("error"))


@app.route("/join", methods=["POST"])
def join():
    name = (request.form.get("team_name") or request.form.get("name") or "").strip()[:60]
    roll = (request.form.get("team_number") or request.form.get("roll") or "").strip().upper()[:30]
    if not name or not roll:
        return redirect("/?error=Enter both Team Name and Team Number")
    if ONLINE.get(roll):
        return redirect("/?error=This Team Number is already logged in on another device")
    pr = state["participants"].setdefault(roll, {"name": name, "warnings": 0})
    pr["name"] = name
    save()
    session["roll"] = roll
    return redirect("/play")


@app.route("/play")
def play():
    if "roll" not in session:
        return redirect("/")
    pr = state["participants"].get(session["roll"], {})
    return render_template("play.html", name=pr.get("name", ""), roll=session["roll"])


@app.route("/admin/login", methods=["GET", "POST"])
def admin_login():
    if request.method == "POST":
        if request.form.get("password") == ADMIN_PASSWORD:
            session["admin"] = True
            return redirect("/admin")
        return render_template("admin_login.html", error="Wrong password")
    return render_template("admin_login.html", error=None)


@app.route("/admin")
def admin():
    if not session.get("admin"):
        return redirect("/admin/login")
    return render_template("admin.html", initial_rounds=ROUNDS)


@app.route("/export.csv")
def export():
    if not session.get("admin"):
        return redirect("/admin/login")
    out = io.StringIO()
    w = csv.writer(out)
    w.writerow(["round", "team_number", "team_name", "prompt", "submitted_at", "times_submitted", "score", "warnings", "status"])
    for key, s in sorted(state["submissions"].items(), key=lambda kv: (kv[1]["round"], kv[1]["time"])):
        pr = state["participants"].get(s["roll"], {})
        status_str = "Disqualified" if pr.get("disqualified") else "Active"
        w.writerow([s["round"] + 1, s["roll"], s["name"], s["prompt"],
                    time.strftime("%H:%M:%S", time.localtime(s["time"])), s["count"],
                    state["scores"].get(key, ""), pr.get("warnings", 0), status_str])
    return Response("\ufeff" + out.getvalue(), mimetype="text/csv",
                    headers={"Content-Disposition": "attachment; filename=promptathon_results.csv"})


@app.route("/export.xlsx")
def export_xlsx():
    if not session.get("admin"):
        return redirect("/admin/login")
    
    wb = openpyxl.Workbook()
    # Sheet 1: Summary & Scores
    ws_summary = wb.active
    ws_summary.title = "Summary & Scores"
    
    header_fill = PatternFill(start_color="14213D", end_color="14213D", fill_type="solid")
    header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
    round_header_fill = PatternFill(start_color="E8920C", end_color="E8920C", fill_type="solid")
    
    headers = ["Team Number", "Team Name", "Overall Score (Out of 145)", "Round 1 (Out of 40)", "Round 2 (Out of 45)", "Round 3 (Out of 60)", "Warnings", "Status"]
    ws_summary.append(headers)
    
    for col in range(1, len(headers) + 1):
        cell = ws_summary.cell(row=1, column=col)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
    
    for roll, pr in sorted(state["participants"].items(), key=lambda x: x[0]):
        r1_sc = float(state["scores"].get(f"{roll}|0", 0) or 0)
        r2_sc = float(state["scores"].get(f"{roll}|1", 0) or 0)
        r3_sc = float(state["scores"].get(f"{roll}|2", 0) or 0)
        tot = round(r1_sc + r2_sc + r3_sc, 1)
        status = "Disqualified" if pr.get("disqualified") else "Active"
        row = [roll, pr["name"], f"{tot} / 145", f"{r1_sc} / 40", f"{r2_sc} / 45", f"{r3_sc} / 60", pr.get("warnings", 0), status]
        ws_summary.append(row)
        
    for r in ws_summary.iter_rows(min_row=2):
        for cell in r:
            cell.alignment = Alignment(horizontal="center", vertical="center")
            
    # Create separate sheets for each Round (without score column)
    for r_idx, r in enumerate(ROUNDS):
        ws_round = wb.create_sheet(title=f"Round {r_idx + 1}")
        r_headers = ["Team Number", "Team Name", "Submitted Prompt (All Themes)", "Submitted At", "Times Submitted", "Warnings", "Status"]
        ws_round.append(r_headers)
        
        for col in range(1, len(r_headers) + 1):
            cell = ws_round.cell(row=1, column=col)
            cell.fill = round_header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")
            
        for key, s in sorted(state["submissions"].items(), key=lambda kv: kv[1]["time"]):
            if s["round"] == r_idx:
                pr = state["participants"].get(s["roll"], {})
                status_str = "Disqualified" if pr.get("disqualified") else "Active"
                ws_round.append([
                    s["roll"],
                    s["name"],
                    s["prompt"],
                    time.strftime("%H:%M:%S", time.localtime(s["time"])),
                    s["count"],
                    pr.get("warnings", 0),
                    status_str
                ])
                
        # Format data rows with wrap_text for prompt column
        for row_cells in ws_round.iter_rows(min_row=2):
            for idx, cell in enumerate(row_cells):
                if idx == 2:  # Submitted Prompt column
                    cell.alignment = Alignment(wrap_text=True, vertical="top")
                else:
                    cell.alignment = Alignment(vertical="top")
                
        # Set specific column widths for neat display
        ws_round.column_dimensions['A'].width = 16
        ws_round.column_dimensions['B'].width = 24
        ws_round.column_dimensions['C'].width = 65
        ws_round.column_dimensions['D'].width = 16
        ws_round.column_dimensions['E'].width = 18
        ws_round.column_dimensions['F'].width = 12
        ws_round.column_dimensions['G'].width = 16
            
    for col in ws_summary.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = openpyxl.utils.get_column_letter(col[0].column)
        ws_summary.column_dimensions[col_letter].width = min(max(max_len + 4, 18), 35)

    out = io.BytesIO()
    wb.save(out)
    out.seek(0)
    return Response(out.read(), mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    headers={"Content-Disposition": "attachment; filename=Promptathon_Results.xlsx"})


# ---------- sockets ----------
@socketio.on("connect")
def on_connect():
    if session.get("admin"):
        ADMIN_SIDS.add(request.sid)
        join_room("admins")
        emit("admin_data", admin_data())
    roll = session.get("roll")
    if roll:
        join_room("participants")
        ONLINE.setdefault(roll, set()).add(request.sid)
        SID_ROLL[request.sid] = roll
        emit("round_state", round_payload())
        emit("me", me_payload(roll))
        push_admin()


@socketio.on("disconnect")
def on_disconnect():
    ADMIN_SIDS.discard(request.sid)
    roll = SID_ROLL.pop(request.sid, None)
    if roll and roll in ONLINE:
        ONLINE[roll].discard(request.sid)
        push_admin()


@socketio.on("violation")
def on_violation(data):
    roll = SID_ROLL.get(request.sid)
    if not roll or state["status"] != "running":
        return
    pr = state["participants"].get(roll)
    if not pr or pr.get("disqualified"):
        return
    kind = str(data.get("type", "unknown"))[:30]
    pr["warnings"] = pr.get("warnings", 0) + 1
    
    # 2 warnings allowed. Reaching out of limit (i.e. 3rd violation / warnings > MAX_WARNINGS) -> Disqualified
    if pr["warnings"] > MAX_WARNINGS:
        pr["disqualified"] = True
        pr["last_event"] = f"DISQUALIFIED ({kind}) at {time.strftime('%H:%M:%S')}"
        pr["disqualified_reason"] = f"Exceeded maximum allowed warnings ({MAX_WARNINGS} max). Disqualified due to {kind}."
        save()
        emit("me", me_payload(roll))
        emit("disqualified", {"reason": pr["disqualified_reason"]})
        push_admin()
        return

    pr["last_event"] = f"{kind} (Warning {pr['warnings']}/{MAX_WARNINGS}) at {time.strftime('%H:%M:%S')}"
    save()
    emit("me", me_payload(roll))
    emit("warning", {"count": pr["warnings"], "max": MAX_WARNINGS, "type": kind})
    push_admin()


@socketio.on("submit")
def on_submit(data):
    roll = SID_ROLL.get(request.sid)
    if not roll:
        return
    pr = state["participants"].get(roll, {})
    if pr.get("disqualified"):
        emit("submit_result", {"ok": False, "msg": "You have been disqualified from participation."})
        return
    if state["status"] != "running" or time.time() >= state["end_time"]:
        emit("submit_result", {"ok": False, "msg": "Round is closed."})
        return
    
    answers = data.get("answers", {})
    text = str(data.get("prompt", "")).strip()[:8000]
    
    if isinstance(answers, dict) and len(answers) > 0:
        cur_round = ROUNDS[state["current"]] if 0 <= state["current"] < len(ROUNDS) else {}
        questions = cur_round.get("questions", [])
        lines = []
        for k, ans in sorted(answers.items(), key=lambda x: int(x[0]) if str(x[0]).isdigit() else 0):
            ans_str = str(ans).strip()
            if ans_str:
                idx = int(k) if str(k).isdigit() else 0
                q_title = questions[idx]["title"] if idx < len(questions) else f"Question {idx+1}"
                lines.append(f"[{q_title}]\n{ans_str}")
        if lines:
            text = "\n\n".join(lines)
            
    if not text:
        emit("submit_result", {"ok": False, "msg": "Please write your prompt before submitting."})
        return
        
    key = f"{roll}|{state['current']}"
    old = state["submissions"].get(key)
    state["submissions"][key] = {
        "roll": roll,
        "name": state["participants"].get(roll, {}).get("name", roll),
        "round": state["current"],
        "prompt": text,
        "answers": answers if isinstance(answers, dict) else {},
        "submitted": True,
        "time": time.time(),
        "count": (old["count"] + 1) if old else 1
    }
    
    # Calculate auto similarity score based on key answers
    auto_score = auto_evaluate_round(state["current"], answers, text)
    state["scores"][key] = auto_score
    
    save()
    emit("submit_result", {
        "ok": True,
        "msg": "Prompt submitted successfully! Please wait for the next round until admin gives access.",
        "prompt": text,
        "answers": answers if isinstance(answers, dict) else {}
    })
    emit("me", me_payload(roll))
    push_admin()



def admin_only(fn):
    def wrapper(*a, **k):
        if request.sid in ADMIN_SIDS:
            return fn(*a, **k)
    return wrapper


@socketio.on("start_round")
@admin_only
def on_start(data):
    i = int(data.get("index", 0))
    if 0 <= i < len(ROUNDS):
        state["current"] = i
        state["status"] = "running"
        r = ROUNDS[i]
        total_mins = r.get("minutes")
        if not total_mins:
            q_list = r.get("questions", [])
            total_mins = sum(q.get("minutes", 15) for q in q_list) or 15
        now = time.time()
        state["start_time"] = now
        state["end_time"] = now + total_mins * 60
        save()
        push()
        for roll in list(ONLINE):
            for sid in ONLINE[roll]:
                socketio.emit("me", me_payload(roll), to=sid)


@socketio.on("end_round")
@admin_only
def on_end(_=None):
    if state["status"] == "running":
        state["status"] = "ended"
        save()
        push()


@socketio.on("add_time")
@admin_only
def on_add(data):
    if state["status"] == "running":
        state["end_time"] += int(data.get("seconds", 60))
        save()
        push()


@socketio.on("set_score")
@admin_only
def on_score(data):
    try:
        state["scores"][data["key"]] = float(data["value"])
    except (ValueError, KeyError):
        state["scores"].pop(data.get("key"), None)
    save()
    push_admin()


@socketio.on("reset_all")
@admin_only
def on_reset(_=None):
    state.update({"current": None, "status": "idle", "end_time": None,
                  "participants": {}, "submissions": {}, "scores": {}})
    save()
    push()


def lan_ip():
    try:
        s = socket.socket(socket.AF_INET, socket.SOCK_DGRAM)
        s.connect(("8.8.8.8", 80))
        return s.getsockname()[0]
    except OSError:
        return "127.0.0.1"


if __name__ == "__main__":
    start_ticker_if_needed()
    port = int(os.environ.get("PORT", 5000))
    ip = lan_ip()
    print(f"\n  Participants open : http://{ip}:{port}/\n  Admin open        : http://{ip}:{port}/admin\n")
    socketio.run(app, host="0.0.0.0", port=port, allow_unsafe_werkzeug=True)
